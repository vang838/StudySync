from dataclasses import dataclass
from typing import BinaryIO, Protocol

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from src.application.docs_upload import DOCX_MEDIA_TYPE, PDF_MEDIA_TYPE, PPTX_MEDIA_TYPE, TEXT_MEDIA_TYPE


class DocumentParseError(Exception):
    pass


class UnsupportedParserError(DocumentParseError):
    pass


@dataclass(frozen=True)
class ExtractedSegment:
    text: str
    source_index: int
    source_label: str


@dataclass(frozen=True)
class ExtractedDocument:
    media_type: str
    segments: list[ExtractedSegment]


class DocumentParser(Protocol):
    def parse(self, fileobj: BinaryIO) -> ExtractedDocument:
        ...


class PdfDocumentParser:
    def parse(self, fileobj: BinaryIO) -> ExtractedDocument:
        try:
            fileobj.seek(0)
            reader = PdfReader(fileobj)

            segments = [
                ExtractedSegment(
                    text=page.extract_text() or "",
                    source_index=index,
                    source_label=f"page:{index}",
                )
                for index, page in enumerate(reader.pages, start=1)
            ]

            fileobj.seek(0)
            return ExtractedDocument(media_type=PDF_MEDIA_TYPE, segments=segments)
        except Exception as exc:
            fileobj.seek(0)
            raise DocumentParseError("Failed to parse PDF document.") from exc


class DocxDocumentParser:
    def parse(self, fileobj: BinaryIO) -> ExtractedDocument:
        try:
            fileobj.seek(0)
            document = Document(fileobj)

            segments = [
                ExtractedSegment(
                    text=paragraph.text,
                    source_index=index,
                    source_label=f"paragraph:{index}",
                )
                for index, paragraph in enumerate(document.paragraphs, start=1)
            ]

            fileobj.seek(0)
            return ExtractedDocument(media_type=DOCX_MEDIA_TYPE, segments=segments)
        except Exception as exc:
            fileobj.seek(0)
            raise DocumentParseError("Failed to parse DOCX document.") from exc


class PptxDocumentParser:
    def parse(self, fileobj: BinaryIO) -> ExtractedDocument:
        try:
            fileobj.seek(0)
            presentation = Presentation(fileobj)
            segments = []

            for slide_index, slide in enumerate(presentation.slides, start=1):
                text_parts = []

                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text_parts.append(shape.text)

                segments.append(
                    ExtractedSegment(
                        text="\n".join(text_parts),
                        source_index=slide_index,
                        source_label=f"slide:{slide_index}",
                    )
                )

            fileobj.seek(0)
            return ExtractedDocument(media_type=PPTX_MEDIA_TYPE, segments=segments)
        except Exception as exc:
            fileobj.seek(0)
            raise DocumentParseError("Failed to parse PPTX document.") from exc


class TextDocumentParser:
    def parse(self, fileobj: BinaryIO) -> ExtractedDocument:
        try:
            fileobj.seek(0)
            text = fileobj.read().decode("utf-8")
            fileobj.seek(0)

            return ExtractedDocument(
                media_type=TEXT_MEDIA_TYPE,
                segments=[
                    ExtractedSegment(
                        text=text,
                        source_index=1,
                        source_label="text:1",
                    )
                ],
            )
        except UnicodeDecodeError as exc:
            fileobj.seek(0)
            raise DocumentParseError("Text document is not valid UTF-8.") from exc

class DocumentParserRegistry:
    def __init__(self):
        self._parsers: dict[str, DocumentParser] = {
            PDF_MEDIA_TYPE: PdfDocumentParser(),
            DOCX_MEDIA_TYPE: DocxDocumentParser(),
            PPTX_MEDIA_TYPE: PptxDocumentParser(),
            TEXT_MEDIA_TYPE: TextDocumentParser(),
        }

    def get(self, media_type: str) -> DocumentParser:
        parser = self._parsers.get(media_type)

        if parser is None:
            raise UnsupportedParserError(f"No parser is available for media type '{media_type}'.")

        return parser