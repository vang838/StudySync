from io import BytesIO
from unittest import TestCase

from docx import Document
from pptx import Presentation

from src.application.document_parser import (
    DocumentParserRegistry,
    UnsupportedParserError,
)
from src.application.docs_upload import DOCX_MEDIA_TYPE, PDF_MEDIA_TYPE, PPTX_MEDIA_TYPE, TEXT_MEDIA_TYPE


class TestDocumentParserRegistry(TestCase):
    def setUp(self):
        self.registry = DocumentParserRegistry()

    def test_text_parser(self):
        parser = self.registry.get(TEXT_MEDIA_TYPE)

        result = parser.parse(BytesIO(b"Operating systems manage processes."))

        self.assertEqual(result.media_type, TEXT_MEDIA_TYPE)
        self.assertEqual(len(result.segments), 1)
        self.assertEqual(result.segments[0].text, "Operating systems manage processes.")
        self.assertEqual(result.segments[0].source_label, "text:1")

    def test_docx_parser(self):
        fileobj = BytesIO()
        document = Document()
        document.add_paragraph("First paragraph")
        document.add_paragraph("Second paragraph")
        document.save(fileobj)
        fileobj.seek(0)

        parser = self.registry.get(DOCX_MEDIA_TYPE)
        result = parser.parse(fileobj)

        self.assertEqual(result.media_type, DOCX_MEDIA_TYPE)
        self.assertEqual(len(result.segments), 2)
        self.assertEqual(result.segments[0].text, "First paragraph")
        self.assertEqual(result.segments[1].source_label, "paragraph:2")

    def test_pptx_parser(self):
        fileobj = BytesIO()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        textbox = slide.shapes.add_textbox(0, 0, 100, 100)
        textbox.text = "CPU Scheduling"
        presentation.save(fileobj)
        fileobj.seek(0)

        parser = self.registry.get(PPTX_MEDIA_TYPE)
        result = parser.parse(fileobj)

        self.assertEqual(result.media_type, PPTX_MEDIA_TYPE)
        self.assertEqual(len(result.segments), 1)
        self.assertIn("CPU Scheduling", result.segments[0].text)
        self.assertEqual(result.segments[0].source_label, "slide:1")

    def test_pdf_parser_is_registered(self):
        parser = self.registry.get(PDF_MEDIA_TYPE)

        self.assertIsNotNone(parser)

    def test_unsupported_media_type(self):
        with self.assertRaises(UnsupportedParserError):
            self.registry.get("application/octet-stream")